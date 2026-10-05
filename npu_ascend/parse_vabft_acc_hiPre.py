#!/usr/bin/env python3
"""Convert vabft_acc_hiPre CSV/JSONL results into XLSX and Markdown reports."""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COUNT_FIELDS = (
    "total_blocks",
    "all_injected",
    "non_nan_injected",
    "nan_injected",
    "inf_injected",
    "finite_injected",
    "non_nan_injected_rows",
    "non_nan_detected_rows",
    "non_nan_injected_blocks",
    "non_nan_detected_blocks",
    "non_nan_localized",
    "nan_injected_rows",
    "nan_detected_rows",
    "nan_injected_blocks",
    "nan_detected_blocks",
    "nan_localized",
    "injected_rows",
    "detected_rows",
    "false_rows",
    "clean_rows",
    "injected_blocks",
    "detected_blocks",
    "false_blocks",
    "clean_blocks",
    "localized",
)

RATE_FIELDS = (
    "non_nan_row_DR",
    "non_nan_block_DR",
    "non_nan_localization_R",
    "non_nan_conditional_localization_R",
    "nan_row_DR",
    "nan_block_DR",
    "nan_localization_R",
    "nan_conditional_localization_R",
    "row_DR",
    "block_DR",
    "row_FPR",
    "block_FPR",
    "localization_R",
    "conditional_localization_R",
    "nan_rate",
    "inf_rate",
)

SUMMARY_FIELDS = (
    "conditions",
    *COUNT_FIELDS,
    *RATE_FIELDS,
)


@dataclass
class AuditSummary:
    records: int = 0
    matrix_lifetimes: int = 0
    duplicate_positions: int = 0
    flips_0_to_1: int = 0
    flips_1_to_0: int = 0
    finite: int = 0
    inf: int = 0
    nan: int = 0


def optional_float(value: str | float | int | None) -> float | None:
    if value is None or value == "":
        return None
    result = float(value)
    return None if math.isnan(result) else result


def load_results(path: Path) -> list[dict[str, object]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        raw_rows = list(csv.DictReader(handle))
    if not raw_rows:
        raise ValueError(f"empty results CSV: {path}")

    required = {
        "K",
        "dtype",
        "distribution",
        "bit",
        "all_injected",
        "nan_injected",
        "non_nan_injected_rows",
        "non_nan_detected_rows",
        "non_nan_localized",
        "nan_injected_rows",
        "nan_detected_rows",
        "nan_localized",
    }
    missing = required - set(raw_rows[0])
    if missing:
        raise ValueError(f"missing CSV columns: {sorted(missing)}")

    rows: list[dict[str, object]] = []
    integer_fields = set(COUNT_FIELDS) | {
        "seed",
        "base_seed",
        "M",
        "N",
        "K",
        "bit",
        "C",
        "refresh_interval",
    }
    float_fields = set(RATE_FIELDS) | {"p_inject", "e_max", "c_sigma"}
    for raw in raw_rows:
        row: dict[str, object] = dict(raw)
        for field in integer_fields:
            if field in row and row[field] not in (None, ""):
                row[field] = int(str(row[field]))
        for field in float_fields:
            if field in row:
                row[field] = optional_float(row[field])
        rows.append(row)
    return rows


def aggregate(rows: Sequence[dict[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {"conditions": len(rows)}
    for field in COUNT_FIELDS:
        result[field] = sum(int(row.get(field, 0) or 0) for row in rows)

    def div(numerator: str, denominator: str) -> float | None:
        den = int(result[denominator])
        return int(result[numerator]) / den if den else None

    result.update(
        {
            "non_nan_row_DR": div(
                "non_nan_detected_rows", "non_nan_injected_rows"
            ),
            "non_nan_block_DR": div(
                "non_nan_detected_blocks", "non_nan_injected_blocks"
            ),
            "non_nan_localization_R": div(
                "non_nan_localized", "non_nan_injected_rows"
            ),
            "non_nan_conditional_localization_R": div(
                "non_nan_localized", "non_nan_detected_rows"
            ),
            "nan_row_DR": div("nan_detected_rows", "nan_injected_rows"),
            "nan_block_DR": div("nan_detected_blocks", "nan_injected_blocks"),
            "nan_localization_R": div("nan_localized", "nan_injected_rows"),
            "nan_conditional_localization_R": div(
                "nan_localized", "nan_detected_rows"
            ),
            "row_DR": div("detected_rows", "injected_rows"),
            "block_DR": div("detected_blocks", "injected_blocks"),
            "row_FPR": div("false_rows", "clean_rows"),
            "block_FPR": div("false_blocks", "clean_blocks"),
            "localization_R": div("localized", "injected_rows"),
            "conditional_localization_R": div("localized", "detected_rows"),
            "nan_rate": div("nan_injected", "all_injected"),
            "inf_rate": div("inf_injected", "all_injected"),
        }
    )
    return result


def group_rows(
    rows: Sequence[dict[str, object]], keys: Sequence[str]
) -> list[dict[str, object]]:
    groups: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[tuple(row[key] for key in keys)].append(row)

    output = []
    for group_key in sorted(groups):
        item = dict(zip(keys, group_key))
        item.update(aggregate(groups[group_key]))
        output.append(item)
    return output


CATEGORY_PREFIX = {"非NaN": "non_nan", "NaN": "nan", "汇总": "total"}


def category_view(summary: dict[str, object], category: str) -> dict[str, object]:
    """Project one aggregate into a common block/row/localization schema."""
    prefix = CATEGORY_PREFIX[category]
    if prefix == "total":
        names = {
            "injected": "all_injected",
            "injected_blocks": "injected_blocks",
            "detected_blocks": "detected_blocks",
            "injected_rows": "injected_rows",
            "detected_rows": "detected_rows",
            "localized": "localized",
            "block_DR": "block_DR",
            "row_DR": "row_DR",
            "localization_R": "localization_R",
            "conditional_localization_R": "conditional_localization_R",
        }
    else:
        names = {
            "injected": f"{prefix}_injected",
            "injected_blocks": f"{prefix}_injected_blocks",
            "detected_blocks": f"{prefix}_detected_blocks",
            "injected_rows": f"{prefix}_injected_rows",
            "detected_rows": f"{prefix}_detected_rows",
            "localized": f"{prefix}_localized",
            "block_DR": f"{prefix}_block_DR",
            "row_DR": f"{prefix}_row_DR",
            "localization_R": f"{prefix}_localization_R",
            "conditional_localization_R": f"{prefix}_conditional_localization_R",
        }
    return {target: summary.get(source) for target, source in names.items()}


def category_summary_rows(
    rows: Sequence[dict[str, object]], category: str
) -> list[dict[str, object]]:
    output = [{"分组": "总体", **category_view(aggregate(rows), category)}]
    dimensions = (
        ("dtype",),
        ("dtype", "bit"),
        ("dtype", "distribution"),
        ("dtype", "K"),
    )
    for keys in dimensions:
        for item in group_rows(rows, keys):
            labels = {key: item[key] for key in keys}
            output.append(
                {
                    "分组": "+".join(keys),
                    **labels,
                    **category_view(item, category),
                }
            )
    return output


def inspect_audit(path: Path | None) -> AuditSummary | None:
    if path is None:
        return None
    summary = AuditSummary()
    positions: dict[tuple[object, ...], set[tuple[int, int]]] = defaultdict(set)
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            record = json.loads(line)
            summary.records += 1
            key = (
                record.get("K"),
                record.get("dtype"),
                record.get("distribution"),
                record.get("bit"),
                record.get("condition_seed"),
                record.get("matrix_generation"),
            )
            position = (int(record["global_row"]), int(record["global_col"]))
            if position in positions[key]:
                summary.duplicate_positions += 1
            positions[key].add(position)

            direction = record.get("direction")
            if direction == "0->1":
                summary.flips_0_to_1 += 1
            elif direction == "1->0":
                summary.flips_1_to_0 += 1
            classification = str(record.get("classification", ""))
            if classification in {"finite", "inf", "nan"}:
                setattr(summary, classification, getattr(summary, classification) + 1)
            else:
                raise ValueError(
                    f"unknown classification at {path}:{line_number}: {classification}"
                )
    summary.matrix_lifetimes = len(positions)
    return summary


def display_name(field: str) -> str:
    names = {
        "conditions": "条件数",
        "分组": "分组",
        "injected": "注错数",
        "all_injected": "全部XOR注错",
        "non_nan_injected": "非NaN注错",
        "nan_injected": "NaN注错",
        "inf_injected": "Inf注错",
        "finite_injected": "有限值注错",
        "non_nan_injected_rows": "非NaN注错行",
        "non_nan_detected_rows": "非NaN行定位",
        "non_nan_injected_blocks": "非NaN注错块",
        "non_nan_detected_blocks": "非NaN块检出",
        "non_nan_localized": "非NaN精确定位",
        "nan_injected_rows": "NaN注错行",
        "nan_detected_rows": "NaN行定位",
        "nan_injected_blocks": "NaN注错块",
        "nan_detected_blocks": "NaN块检出",
        "nan_localized": "NaN精确定位",
        "injected_rows": "注错行",
        "detected_rows": "检出行",
        "false_rows": "误检行",
        "clean_rows": "无错行",
        "injected_blocks": "注错块",
        "detected_blocks": "检出块",
        "false_blocks": "误检块",
        "clean_blocks": "无错块",
        "localized": "精确定位",
        "non_nan_row_DR": "非NaN行定位率",
        "non_nan_block_DR": "非NaN块检出率",
        "non_nan_localization_R": "非NaN精确定位率",
        "non_nan_conditional_localization_R": "非NaN检出后定位率",
        "nan_row_DR": "NaN行定位率",
        "nan_block_DR": "NaN块检出率",
        "nan_localization_R": "NaN精确定位率",
        "nan_conditional_localization_R": "NaN检出后定位率",
        "row_DR": "行检出率",
        "block_DR": "块检出率",
        "row_FPR": "行误检率",
        "block_FPR": "块误检率",
        "localization_R": "定位率",
        "conditional_localization_R": "检出后定位率",
        "nan_rate": "NaN率",
        "inf_rate": "Inf率",
    }
    return names.get(field, field)


def write_sheet(
    workbook: Workbook,
    title: str,
    rows: Sequence[dict[str, object]],
    columns: Sequence[str],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([display_name(column) for column in columns])
    for row in rows:
        sheet.append([row.get(column) for column in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for index, column in enumerate(columns, 1):
        if column in RATE_FIELDS:
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = "0.00%"
        values = [str(sheet.cell(row=row, column=index).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(max(map(len, values)) + 2, 10), 32
        )


def create_xlsx(
    path: Path,
    rows: Sequence[dict[str, object]],
    audit: AuditSummary | None,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    detail_columns = list(rows[0])
    write_sheet(workbook, "完整数据", rows, detail_columns)
    category_columns = (
        "分组", "dtype", "bit", "distribution", "K", "injected",
        "injected_blocks", "detected_blocks", "block_DR", "injected_rows",
        "detected_rows", "row_DR", "localized", "localization_R",
        "conditional_localization_R",
    )
    for category in CATEGORY_PREFIX:
        write_sheet(
            workbook,
            f"{category}汇总" if category != "汇总" else "总体汇总",
            category_summary_rows(rows, category),
            category_columns,
        )
    if audit is not None:
        audit_row = {
            "records": audit.records,
            "matrix_lifetimes": audit.matrix_lifetimes,
            "duplicate_positions": audit.duplicate_positions,
            "flips_0_to_1": audit.flips_0_to_1,
            "flips_1_to_0": audit.flips_1_to_0,
            "finite": audit.finite,
            "inf": audit.inf,
            "nan": audit.nan,
        }
        write_sheet(workbook, "审计汇总", [audit_row], tuple(audit_row))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def percent(value: object) -> str:
    return "NA" if value is None else f"{float(value) * 100:.2f}%"


def markdown_table(
    rows: Sequence[dict[str, object]], columns: Sequence[str], rate_fields: set[str]
) -> list[str]:
    output = [
        "| " + " | ".join(display_name(column) for column in columns) + " |",
        "|" + "|".join("---" for _ in columns) + "|",
    ]
    for row in rows:
        values = []
        for column in columns:
            value = row.get(column)
            if column in rate_fields:
                values.append(percent(value))
            elif value is None:
                values.append("NA")
            else:
                values.append(str(value))
        output.append("| " + " | ".join(values) + " |")
    return output


def create_markdown(
    path: Path,
    rows: Sequence[dict[str, object]],
    source_csv: Path,
    audit_path: Path | None,
    audit: AuditSummary | None,
) -> None:
    first = rows[0]
    lines = [
        "# V-ABFT 错误注入实验报告",
        "",
        "## 1. 实验信息",
        "",
        f"- 数据源：`{source_csv}`",
        f"- run ID：`{first.get('run_id', 'unknown')}`",
        f"- 时间：`{first.get('timestamp', 'unknown')}`",
        f"- 基础随机种子：`{first.get('base_seed', 'unknown')}`",
        f"- 形状：`M={first.get('M')}, N={first.get('N')}`",
        f"- K：`{sorted({int(row['K']) for row in rows})}`",
        f"- 每条件 trial：`{first.get('C')}`",
        f"- 矩阵刷新间隔：`{first.get('refresh_interval')}`",
        f"- block 注错概率：`{first.get('p_inject')}`",
        f"- 条件数：`{len(rows)}`",
        "",
    ]
    metric_columns = (
        "injected", "injected_blocks", "detected_blocks", "block_DR",
        "injected_rows", "detected_rows", "row_DR", "localized",
        "localization_R", "conditional_localization_R",
    )
    dimension_specs = (
        ("按 dtype", ("dtype",)),
        ("按指数位", ("dtype", "bit")),
        ("按输入分布", ("dtype", "distribution")),
        ("按 K", ("dtype", "K")),
    )
    for section_number, category in enumerate(CATEGORY_PREFIX, 2):
        lines.extend([f"## {section_number}. {category}", "", "### 总体", ""])
        overall = category_view(aggregate(rows), category)
        lines.extend(markdown_table([overall], metric_columns, set(RATE_FIELDS)))
        for title, keys in dimension_specs:
            projected = []
            for item in group_rows(rows, keys):
                projected.append(
                    {
                        **{key: item[key] for key in keys},
                        **category_view(item, category),
                    }
                )
            lines.extend(["", f"### {title}", ""])
            lines.extend(
                markdown_table(
                    projected, (*keys, *metric_columns), set(RATE_FIELDS)
                )
            )

    lines.extend(["", "## 5. 注错审计", ""])
    if audit is None:
        lines.append("未提供 audit JSONL，未执行坐标去重与翻转方向检查。")
    else:
        lines.extend(
            [
                f"- 审计文件：`{audit_path}`",
                f"- 注错记录：{audit.records}",
                f"- 矩阵生命周期：{audit.matrix_lifetimes}",
                f"- 同一矩阵内重复坐标：{audit.duplicate_positions}",
                f"- `0→1`：{audit.flips_0_to_1}",
                f"- `1→0`：{audit.flips_1_to_0}",
                f"- finite / Inf / NaN：{audit.finite} / {audit.inf} / {audit.nan}",
            ]
        )
    lines.extend(
        [
            "",
            "## 6. 说明",
            "",
            "- 所有汇总率均由整数计数加总后重新计算，不平均条件级百分比。",
            "- 统一检出规则为 D1 差值是 NaN/Inf，或 `|D1| >= threshold`。",
            "- NaN 仅单列统计；D1、D2 均为 NaN，无法通过 D2/D1 定位列，因此精确定位记为失败。",
            "- 汇总项的分母包含非 NaN 与 NaN，且汇总计数等于两类计数之和。",
            "- `NA` 表示分母为零，不能解释为 0%。",
        ]
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="summary CSV")
    parser.add_argument("--audit", type=Path, help="optional injection audit JSONL")
    parser.add_argument("--xlsx", type=Path, required=True, help="output workbook")
    parser.add_argument("--markdown", type=Path, required=True, help="output report")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    rows = load_results(args.input)
    audit = inspect_audit(args.audit)
    create_xlsx(args.xlsx, rows, audit)
    create_markdown(args.markdown, rows, args.input, args.audit, audit)
    print(f"Loaded {len(rows)} conditions from {args.input}")
    print(f"XLSX saved to {args.xlsx}")
    print(f"Markdown saved to {args.markdown}")
    if audit is not None:
        print(
            f"Audit: {audit.records} records, "
            f"{audit.matrix_lifetimes} matrix lifetimes, "
            f"{audit.duplicate_positions} duplicate positions"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
